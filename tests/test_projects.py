from __future__ import annotations

import sqlite3

from openpyxl import load_workbook

from src.database import (
    add_project_record,
    create_project,
    delete_project_record,
    initialize_database,
    load_cell_styles,
    load_column_display_names,
    list_module_display_names,
    list_projects,
    load_project,
    rename_module_display,
    rename_project,
    save_cell_styles,
    save_column_display_names,
    save_project,
)
from src.data_models import DEFAULT_PROJECT_NAME, TABLE_SCHEMAS
from src.excel_exporter import export_project_to_excel
from src.sample_data import sample_project_data


def test_projects_table_created_and_default_project_added(tmp_path):
    project = tmp_path / "sample.pfcproj"
    create_project(project)

    projects = list_projects(project)

    assert projects[0]["project_name"] == DEFAULT_PROJECT_NAME
    with sqlite3.connect(project) as connection:
        assert connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='projects'").fetchone()
        assert connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='module_display_names'").fetchone()
        assert connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='column_display_names'").fetchone()


def test_module_display_names_save_and_load(tmp_path):
    project = tmp_path / "sample.pfcproj"
    create_project(project)
    project_id = list_projects(project)[0]["id"]

    rename_module_display(project, project_id, "electrical_spec", "전기 특성표")
    reloaded = list_module_display_names(project, project_id)

    assert reloaded["electrical_spec"] == "전기 특성표"


def test_project_rename_persists_after_reopen(tmp_path):
    project = tmp_path / "sample.pfcproj"
    create_project(project)
    project_id = list_projects(project)[0]["id"]

    rename_project(project, project_id, "Custom PFC Project")
    initialize_database(project)
    reloaded = list_projects(project)

    assert reloaded[0]["project_name"] == "Custom PFC Project"


def test_delete_project_removes_project_metadata_and_rows(tmp_path):
    project = tmp_path / "sample.pfcproj"
    create_project(project, sample=True)
    custom_id = add_project_record(project, "Delete Me")
    custom_data = sample_project_data()
    custom_data["drawing_request_summary"].loc[0, "Product"] = "Delete Product"
    save_project(project, custom_data, custom_id)
    save_cell_styles(project, {"drawing_request_summary": {(0, "Product"): "#FFFF00"}}, custom_id)

    delete_project_record(project, custom_id)

    projects = list_projects(project)
    deleted_data = load_project(project, custom_id)
    with sqlite3.connect(project) as connection:
        module_rows = connection.execute(
            "SELECT COUNT(*) FROM module_display_names WHERE project_id = ?",
            (str(custom_id),),
        ).fetchone()[0]
        data_rows = connection.execute(
            'SELECT COUNT(*) FROM drawing_request_summary WHERE "_project_id" = ?',
            (str(custom_id),),
        ).fetchone()[0]
        style_rows = connection.execute(
            "SELECT COUNT(*) FROM cell_styles WHERE project_id = ?",
            (str(custom_id),),
        ).fetchone()[0]

    assert all(str(row["id"]) != str(custom_id) for row in projects)
    assert module_rows == 0
    assert data_rows == 0
    assert style_rows == 0
    assert deleted_data["drawing_request_summary"].empty


def test_cell_styles_save_and_load(tmp_path):
    project = tmp_path / "sample.pfcproj"
    create_project(project)
    project_id = list_projects(project)[0]["id"]

    save_cell_styles(
        project,
        {
            "drawing_request_summary": {
                (0, "Project"): "#FFFF00",
                (1, "Product"): "#DDEEFF",
            }
        },
        project_id,
    )
    loaded = load_cell_styles(project, project_id)

    assert loaded["drawing_request_summary"][(0, "Project")] == "#FFFF00"
    assert loaded["drawing_request_summary"][(1, "Product")] == "#DDEEFF"


def test_column_display_names_save_and_load(tmp_path):
    project = tmp_path / "sample.pfcproj"
    create_project(project)
    project_id = list_projects(project)[0]["id"]

    save_column_display_names(
        project,
        {"esms": {"Customer Spec": "고객 사양", "Clause": "조항"}},
        project_id,
    )
    loaded = load_column_display_names(project, project_id)

    assert loaded["esms"]["Customer Spec"] == "고객 사양"
    assert loaded["esms"]["Clause"] == "조항"


def test_legacy_sample_without_project_info_gets_default_project(tmp_path):
    project = tmp_path / "legacy.pfcproj"
    with sqlite3.connect(project) as connection:
        columns = TABLE_SCHEMAS["electrical_spec"]
        connection.execute(
            "CREATE TABLE electrical_spec (" + ", ".join(f'"{column}" TEXT' for column in columns) + ")"
        )
        placeholders = ", ".join(["?"] * len(columns))
        connection.execute(
            "INSERT INTO electrical_spec VALUES (" + placeholders + ")",
            ["1", "Inductance", "L", "", "", "530", "", "uH", "L=530uH", "Y", "", "", "High", ""],
        )
        connection.commit()

    initialize_database(project)
    projects = list_projects(project)
    data = load_project(project, projects[0]["id"])

    assert projects[0]["project_name"] == DEFAULT_PROJECT_NAME
    assert data["electrical_spec"].loc[0, "Item"] == "Inductance"


def test_project_specific_export_uses_selected_project(tmp_path):
    project = tmp_path / "sample.pfcproj"
    output = tmp_path / "custom.xlsx"
    create_project(project, sample=True)
    custom_id = add_project_record(project, "Custom Project")
    custom_data = sample_project_data()
    custom_data["drawing_request_summary"].loc[0, "Product"] = "Custom Product"
    save_project(project, custom_data, custom_id)

    export_project_to_excel(project, output, project_id=custom_id)
    workbook = load_workbook(output)

    assert workbook["Drawing_Request_Summary"]["A2"].value == "Custom Project"
    assert workbook["Drawing_Request_Summary"]["B2"].value == "Custom Product"
